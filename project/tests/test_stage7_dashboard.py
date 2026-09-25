"""Automated tests for Stage 7 — Web Dashboard & Real-Time Project Excel Engine."""
import sys
from pathlib import Path

PROJECT_DIR = Path(__file__).resolve().parent.parent
if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

import io
import json
import unittest
from unittest.mock import patch, MagicMock
from datetime import date, datetime

import openpyxl
from fastapi.testclient import TestClient

from api.index import app
from services.prediction import _logistic_s_curve, calculate_project_prediction
from services.excel_service import (
    get_project_fields,
    get_project_spreadsheet_data,
    generate_project_excel,
    sync_excel_row_to_supabase,
    create_new_project,
    import_project_excel,
)


class TestStage7Dashboard(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        self.sample_project_id = "e42d5bb4-117f-4c63-99a9-8a013129e002"
        self.sample_project = {
            "id": self.sample_project_id,
            "project_name": "Metro Line Extension",
            "project_code": "MLE-01",
            "location": "Pune",
            "status": "active",
            "created_at": datetime(2026, 9, 1, 10, 0, 0),
        }
        self.sample_records = [
            {
                "id": "11111111-1111-1111-1111-111111111111",
                "project_id": self.sample_project_id,
                "record_type": "daily_log",
                "record_date": date(2026, 9, 2),
                "title": "Site clearance",
                "description": "Cleared sector 1",
                "amount": None,
                "unit": None,
                "data": {"workers_count": 15, "progress": 10.0},
                "created_at": datetime(2026, 9, 2, 10, 0, 0),
                "updated_at": datetime(2026, 9, 2, 10, 0, 0),
            },
            {
                "id": "22222222-2222-2222-2222-222222222222",
                "project_id": self.sample_project_id,
                "record_type": "expense",
                "record_date": date(2026, 9, 5),
                "title": "Fuel",
                "description": "Diesel for generator",
                "amount": 3500.0,
                "unit": "INR",
                "data": {"category": "Fuel"},
                "created_at": datetime(2026, 9, 5, 12, 0, 0),
                "updated_at": datetime(2026, 9, 5, 12, 0, 0),
            },
        ]
        self.sample_custom_fields = [
            {
                "record_type": "daily_log",
                "field_name": "workers_count",
                "field_type": "integer",
                "required": False,
                "default_value": None,
                "validation_rules": {},
            },
            {
                "record_type": "expense",
                "field_name": "category",
                "field_type": "text",
                "required": True,
                "default_value": None,
                "validation_rules": {},
            },
        ]

    # 1. Prediction Algorithm Tests
    def test_logistic_s_curve(self):
        self.assertEqual(_logistic_s_curve(0.0), 0.0)
        self.assertEqual(_logistic_s_curve(1.0), 100.0)
        mid_val = _logistic_s_curve(0.5)
        self.assertTrue(45.0 <= mid_val <= 55.0)

    @patch("services.prediction.execute_query")
    def test_calculate_project_prediction(self, mock_query):
        def query_side_effect(sql, params=None):
            sql_clean = " ".join(sql.split()).upper()
            if "FROM PUBLIC.PROJECTS" in sql_clean:
                return [self.sample_project]
            if "FROM PUBLIC.PROJECT_RECORDS" in sql_clean:
                return self.sample_records
            return []

        mock_query.side_effect = query_side_effect

        result = calculate_project_prediction(
            self.sample_project_id,
            target_duration_days=60,
            reference_date=date(2026, 9, 15),
        )

        self.assertEqual(result["project_id"], self.sample_project_id)
        self.assertEqual(result["project_name"], "Metro Line Extension")
        self.assertTrue(len(result["dates"]) > 0)
        self.assertEqual(len(result["dates"]), len(result["expected_work"]))
        self.assertEqual(len(result["dates"]), len(result["actual_work"]))
        self.assertEqual(len(result["dates"]), len(result["predicted_work"]))

        metrics = result["metrics"]
        self.assertIn("planned_completion_date", metrics)
        self.assertIn("predicted_completion_date", metrics)
        self.assertIn("current_progress_pct", metrics)
        self.assertIn("schedule_status", metrics)
        self.assertIn(metrics["schedule_status"], ["Ahead of Schedule", "On Track", "Behind Schedule"])

    # 2. Dedicated Excel Spreadsheet & Schema Tests
    @patch("services.excel_service.execute_query")
    def test_get_project_spreadsheet_data(self, mock_query):
        def query_side_effect(sql, params=None):
            sql_clean = " ".join(sql.split()).upper()
            if "FROM PUBLIC.PROJECTS" in sql_clean:
                return [self.sample_project]
            if "FROM PUBLIC.PROJECT_FIELD_DEFINITIONS" in sql_clean:
                return self.sample_custom_fields
            if "FROM PUBLIC.PROJECT_RECORDS" in sql_clean:
                return self.sample_records
            return []

        mock_query.side_effect = query_side_effect

        data = get_project_spreadsheet_data(self.sample_project_id)
        self.assertEqual(data["project"]["project_name"], "Metro Line Extension")
        self.assertIn("daily_log", data["sheets"])
        self.assertIn("expense", data["sheets"])
        self.assertIn("equipment_log", data["sheets"])

        # Check that custom field category is included in expense sheet
        expense_cols = [c["name"] for c in data["sheets"]["expense"]["columns"]]
        self.assertIn("amount", expense_cols)
        self.assertIn("category", expense_cols)

        # Check rows unpacked
        expense_rows = data["sheets"]["expense"]["rows"]
        self.assertEqual(len(expense_rows), 1)
        self.assertEqual(expense_rows[0]["title"], "Fuel")
        self.assertEqual(expense_rows[0]["amount"], 3500.0)
        self.assertEqual(expense_rows[0]["category"], "Fuel")

    # 3. Excel Binary Generation (.xlsx) Tests
    @patch("services.excel_service.execute_query")
    def test_generate_project_excel(self, mock_query):
        def query_side_effect(sql, params=None):
            sql_clean = " ".join(sql.split()).upper()
            if "FROM PUBLIC.PROJECTS" in sql_clean:
                return [self.sample_project]
            if "FROM PUBLIC.PROJECT_FIELD_DEFINITIONS" in sql_clean:
                return self.sample_custom_fields
            if "FROM PUBLIC.PROJECT_RECORDS" in sql_clean:
                return self.sample_records
            return []

        mock_query.side_effect = query_side_effect

        excel_bytes = generate_project_excel(self.sample_project_id)
        self.assertIsInstance(excel_bytes, bytes)
        self.assertTrue(len(excel_bytes) > 1000)

        # Verify workbook can be read by openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames
        self.assertIn("Daily Logs", sheet_names)
        self.assertIn("Expenses", sheet_names)
        self.assertIn("Equipment Logs", sheet_names)

        # Verify Expenses sheet content
        ws = wb["Expenses"]
        # Header row is at row 3
        headers = [ws.cell(row=3, column=c).value for c in range(1, 8) if ws.cell(row=3, column=c).value]
        self.assertIn("Amount", headers)
        self.assertIn("Category", headers)

    # 4. Excel Row Sync to Supabase Tests
    @patch("services.excel_service.get_db_connection")
    @patch("services.excel_service.execute_query")
    def test_sync_excel_row_insert(self, mock_query, mock_get_conn):
        def query_side_effect(sql, params=None):
            sql_clean = " ".join(sql.split()).upper()
            if "FROM PUBLIC.PROJECTS" in sql_clean:
                return [self.sample_project]
            if "FROM PUBLIC.PROJECT_FIELD_DEFINITIONS" in sql_clean:
                return self.sample_custom_fields
            return []

        mock_query.side_effect = query_side_effect

        mock_conn = MagicMock()
        mock_cur = MagicMock()
        mock_get_conn.return_value.__enter__.return_value = mock_conn
        mock_conn.cursor.return_value.__enter__.return_value = mock_cur

        mock_cur.fetchone.return_value = {
            "id": "99999999-9999-9999-9999-999999999999",
            "project_id": self.sample_project_id,
            "record_type": "expense",
            "record_date": date(2026, 9, 20),
            "title": "Site Hardware",
            "description": "Nails and screws",
            "amount": 1200.0,
            "unit": "INR",
            "data": {"category": "Hardware"},
        }

        row_payload = {
            "record_type": "expense",
            "record_date": "2026-09-20",
            "title": "Site Hardware",
            "amount": 1200.0,
            "unit": "INR",
            "category": "Hardware",
            "description": "Nails and screws",
        }

        res = sync_excel_row_to_supabase(self.sample_project_id, row_payload)
        self.assertTrue(res["success"])
        self.assertEqual(res["record"]["id"], "99999999-9999-9999-9999-999999999999")
        self.assertEqual(res["record"]["title"], "Site Hardware")

    @patch("services.excel_service.execute_query")
    def test_sync_excel_row_missing_required(self, mock_query):
        def query_side_effect(sql, params=None):
            sql_clean = " ".join(sql.split()).upper()
            if "FROM PUBLIC.PROJECTS" in sql_clean:
                return [self.sample_project]
            if "FROM PUBLIC.PROJECT_FIELD_DEFINITIONS" in sql_clean:
                return self.sample_custom_fields
            return []

        mock_query.side_effect = query_side_effect

        # Missing required custom field 'category'
        row_payload = {
            "record_type": "expense",
            "record_date": "2026-09-20",
            "title": "Site Hardware",
            "amount": 1200.0,
        }

        res = sync_excel_row_to_supabase(self.sample_project_id, row_payload)
        self.assertFalse(res["success"])
        self.assertIn("Category", res["error"])

    # 5. FastAPI HTTP Endpoints Tests
    def test_dashboard_html_route(self):
        resp = self.client.get("/dashboard")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("text/html", resp.headers["content-type"])
        self.assertIn("Construction CRM", resp.text)
        self.assertIn("Work Prediction vs. Date", resp.text)
        self.assertIn("Dedicated Project Excel", resp.text)

    @patch("services.database.execute_query")
    def test_api_projects_list(self, mock_query):
        mock_query.return_value = [self.sample_project]
        resp = self.client.get("/api/projects")
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertIn("projects", data)
        self.assertEqual(len(data["projects"]), 1)
        self.assertEqual(data["projects"][0]["project_name"], "Metro Line Extension")

    @patch("services.prediction.calculate_project_prediction")
    def test_api_prediction_route(self, mock_calc):
        mock_calc.return_value = {
            "project_id": self.sample_project_id,
            "dates": ["2026-09-01", "2026-09-02"],
            "expected_work": [0.0, 5.0],
            "actual_work": [0.0, 10.0],
            "predicted_work": [None, 10.0],
            "metrics": {"schedule_status": "Ahead of Schedule"},
        }
        resp = self.client.get(f"/api/projects/{self.sample_project_id}/prediction")
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data["metrics"]["schedule_status"], "Ahead of Schedule")

    @patch("services.excel_service.generate_project_excel")
    @patch("services.excel_service.get_project_spreadsheet_data")
    def test_api_download_excel_route(self, mock_data, mock_gen):
        mock_data.return_value = {"project": {"project_code": "MLE-01"}}
        mock_gen.return_value = b"fake-excel-binary-content"

        resp = self.client.get(f"/api/projects/{self.sample_project_id}/excel")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", resp.headers["content-type"])
        self.assertIn("MLE-01_records.xlsx", resp.headers["content-disposition"])
        self.assertEqual(resp.content, b"fake-excel-binary-content")

    # 6. Project Creation & Excel Import Tests
    @patch("services.excel_service.get_db_connection")
    def test_create_new_project_service(self, mock_get_conn):
        mock_cur = MagicMock()
        mock_conn = MagicMock()
        mock_conn.__enter__.return_value = mock_conn
        mock_conn.cursor.return_value.__enter__.return_value = mock_cur
        mock_get_conn.return_value = mock_conn
        mock_cur.fetchone.return_value = {
            "id": "33333333-3333-3333-3333-333333333333",
            "project_name": "Skyline Tower",
            "project_code": "SKT-01",
            "location": "Mumbai",
            "status": "active",
        }

        res = create_new_project("Skyline Tower", "SKT-01", "Mumbai", "active")
        self.assertTrue(res["success"])
        self.assertEqual(res["project"]["project_name"], "Skyline Tower")
        self.assertGreater(mock_cur.execute.call_count, 1)  # created project + seeded field definitions

    @patch("services.excel_service.execute_query")
    @patch("services.excel_service.get_db_connection")
    def test_import_project_excel_service_reconciliation(self, mock_get_conn, mock_query):
        # Create an Excel workbook in memory
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Daily Logs"
        ws["A1"] = "Metro Line Extension (MLE-01) — Daily Logs"
        ws.append(["Record ID", "Date", "Title", "Workers Count", "Soil Density", "Notes"])
        # Updating existing record
        ws.append(["11111111-1111-1111-1111-111111111111", "2026-09-02", "Updated Site Clearance", "25", "1.85 g/cm3", "Compacted properly"])
        # New row without ID
        ws.append(["", "2026-09-03", "Piling work", "18", "1.90 g/cm3", "Piling complete"])
        out = io.BytesIO()
        wb.save(out)
        excel_bytes = out.getvalue()

        def query_side_effect(sql, params=None):
            sql_clean = " ".join(sql.split()).upper()
            if "FROM PUBLIC.PROJECTS" in sql_clean:
                return [self.sample_project]
            if "FROM PUBLIC.PROJECT_FIELD_DEFINITIONS" in sql_clean:
                return self.sample_custom_fields
            return []

        mock_query.side_effect = query_side_effect

        mock_cur = MagicMock()
        mock_conn = MagicMock()
        mock_conn.__enter__.return_value = mock_conn
        mock_conn.cursor.return_value.__enter__.return_value = mock_cur
        mock_get_conn.return_value = mock_conn

        last_query = [""]

        def execute_side_effect(sql, params=None):
            last_query[0] = " ".join(sql.split()).upper()

        def fetchone_side_effect():
            q = last_query[0]
            if "FROM PUBLIC.PROJECTS WHERE ID" in q:
                return {
                    "id": self.sample_project_id,
                    "project_name": "Metro Line Extension",
                    "project_code": "MLE-01",
                }
            if "FROM PUBLIC.PROJECT_RECORDS WHERE ID" in q:
                return {"id": "11111111-1111-1111-1111-111111111111"}
            if "INSERT INTO PUBLIC.PROJECT_RECORDS" in q:
                return {"id": "88888888-8888-8888-8888-888888888888"}
            return None

        def fetchall_side_effect():
            q = last_query[0]
            if "FROM PUBLIC.PROJECT_FIELD_DEFINITIONS" in q:
                return [
                    {"field_name": "workers_count", "field_type": "integer"}
                ]
            if "SELECT ID FROM PUBLIC.PROJECT_RECORDS" in q:
                return [
                    {"id": "11111111-1111-1111-1111-111111111111"},
                    {"id": "99999999-9999-9999-9999-999999999999"},
                ]
            return []

        mock_cur.execute.side_effect = execute_side_effect
        mock_cur.fetchone.side_effect = fetchone_side_effect
        mock_cur.fetchall.side_effect = fetchall_side_effect

        res = import_project_excel(excel_bytes, self.sample_project_id)
        self.assertTrue(res["success"])
        rec = res["reconciliation"]
        self.assertEqual(rec["updated"], 1)
        self.assertEqual(rec["inserted"], 1)
        self.assertEqual(rec["deleted"], 1)  # 99999999-... was omitted and deleted
        self.assertEqual(rec["custom_fields_added"], 1)  # soil_density registered

    @patch("services.excel_service.create_new_project")
    def test_api_create_project_route(self, mock_create):
        mock_create.return_value = {
            "success": True,
            "project": {"id": "123", "project_name": "Expressway Ph 1"},
        }
        resp = self.client.post(
            "/api/projects/create",
            json={"project_name": "Expressway Ph 1", "project_code": "EXP-01"},
        )
        self.assertEqual(resp.status_code, 201)
        data = resp.json()
        self.assertTrue(data["success"])
        self.assertEqual(data["project"]["project_name"], "Expressway Ph 1")

    @patch("services.excel_service.import_project_excel")
    def test_api_import_excel_route(self, mock_import):
        mock_import.return_value = {
            "success": True,
            "project_id": self.sample_project_id,
            "reconciliation": {"inserted": 2, "updated": 3, "deleted": 1, "custom_fields_added": 0},
        }
        fake_file = io.BytesIO(b"fake-xlsx-bytes")
        resp = self.client.post(
            f"/api/projects/{self.sample_project_id}/import_excel",
            files={"file": ("test.xlsx", fake_file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data["success"])
        self.assertEqual(data["reconciliation"]["inserted"], 2)


if __name__ == "__main__":
    unittest.main()
